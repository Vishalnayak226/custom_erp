#!/usr/bin/env python3
"""Builds the Bulk Purchase Upload template (Stage 129.1): the Excel file a
client fills to make POs and goods receipts in one upload
(internal/server/bulk_purchase.go).

This is a dev-time generator only - never imported or run by the Go binary,
so it adds no runtime dependency. To change the template:

    py scripts/gen_bulk_purchase_template.py %TEMP%/bulk_purchase_template.xlsx
    py scripts/gen_bulk_purchase_template.py %TEMP%/Client_Bulk_Purchase.xlsx CLIENT

then open the file in Excel and save it (so the formula results are stored),
and copy the first one over internal/server/bulk_purchase_template.xlsx (the
blank template the app serves). The second is a copy titled for one client.
Requires openpyxl (pip install openpyxl); not vendored into go.mod/go.sum.

Layout follows the Setup Workbook convention: row 1 = headers, row 2 = grey
hint row whose first cell starts with "#", data from row 3. Money headers
carry "(₹)" so the server's header reader (which drops a bracketed note and a
trailing *) sees plain names.
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
CLIENT = sys.argv[2] if len(sys.argv) > 2 else ""   # "MINN" -> the title says so
N_ROWS = 1000          # Purchase data rows
N_BILLS = 300          # Bills data rows
N_VENDORS = 200        # Vendors data rows
FIRST = 3              # first data row

FONT = "Arial"
C_REQ = "FFF2CC"       # required input
C_OPT = "FFFFFF"       # optional input
C_GRN = "E2EFDA"       # goods-received input
C_AUTO = "EDEDED"      # automatic
C_RED, C_RED_TXT = "FFC7CE", "9C0006"
H_PO, H_PROD, H_TAX, H_QTY, H_GRN, H_AUTO = "1F3864", "1F5C66", "5B3A7A", "7F4F00", "375623", "595959"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    return Font(**kw)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


# ---- Fixed reference lists (sheet "Codes", protected) --------------------------
# HSN: Customs Tariff chapter 71. Every one is GST 3%.
HSN = [
    ("7117 - Imitation / fashion jewellery", 3, "Artificial jewellery of any material. Use this when unsure."),
    ("711711 - Imitation cufflinks and studs", 3, "Cufflinks and studs of base metal, plated or not."),
    ("711719 - Imitation jewellery of base metal", 3, "Artificial, AD / zircon, oxidised, kundan, gold-plated jewellery of brass, copper, alloy, German silver."),
    ("711790 - Imitation jewellery of other materials", 3, "Beads, glass, lac, thread, plastic..."),
    ("7113 - Jewellery of precious metal", 3, "Real gold, silver or platinum jewellery."),
    ("711311 - Silver jewellery", 3, "Of silver, plated or clad with another precious metal or not."),
    ("711319 - Gold, platinum and diamond jewellery", 3, "Of gold, platinum or other precious metal (not silver)."),
    ("711320 - Base metal clad with precious metal", 3, "Rolled gold: a layer of gold or silver bonded on, not just plated."),
    ("711411 - Silverware", 3, "Silver utensils, idols and pooja articles."),
    ("711419 - Gold articles", 3, "Gold utensils, idols and pooja articles."),
    ("711610 - Pearl jewellery and articles", 3, "Natural or cultured pearls."),
    ("711620 - Gemstone jewellery and articles", 3, "Precious and semi-precious stones."),
]
# GST state codes (the first two digits of a GSTIN).
STATES = [
    ("Andaman and Nicobar Islands", "35"), ("Andhra Pradesh", "37"), ("Arunachal Pradesh", "12"),
    ("Assam", "18"), ("Bihar", "10"), ("Chandigarh", "04"), ("Chhattisgarh", "22"),
    ("Dadra and Nagar Haveli and Daman and Diu", "26"), ("Delhi", "07"), ("Goa", "30"), ("Gujarat", "24"),
    ("Haryana", "06"), ("Himachal Pradesh", "02"), ("Jammu and Kashmir", "01"), ("Jharkhand", "20"),
    ("Karnataka", "29"), ("Kerala", "32"), ("Ladakh", "38"), ("Lakshadweep", "31"), ("Madhya Pradesh", "23"),
    ("Maharashtra", "27"), ("Manipur", "14"), ("Meghalaya", "17"), ("Mizoram", "15"), ("Nagaland", "13"),
    ("Odisha", "21"), ("Puducherry", "34"), ("Punjab", "03"), ("Rajasthan", "08"), ("Sikkim", "11"),
    ("Tamil Nadu", "33"), ("Telangana", "36"), ("Tripura", "16"), ("Uttar Pradesh", "09"),
    ("Uttarakhand", "05"), ("West Bengal", "19"),
]
HSN_CODES = f"Codes!$A$2:$A${1 + len(HSN)}"
HSN_GST = f"Codes!$B$2:$B${1 + len(HSN)}"
STATE_NAMES = f"Codes!$E$2:$E${1 + len(STATES)}"
STATE_CODES = f"Codes!$F$2:$F${1 + len(STATES)}"

# ---- Lists (drop-down values; samples the client replaces) -----------------
LISTS = [
    ("Store", ["Main Store"]),
    ("Product Type", ["Earring", "Necklace", "Bracelet", "Ring", "Bangle", "Keychain"]),
    ("Category", ["Antique", "Kundan", "Temple", "CZ", "Trendy", "Imported", "Imported CZ"]),
    ("Style", ["Hoop", "Huggie Hoop", "Double Line Hoop", "Ball Drop", "Flower Stud", "Star Stud",
               "Beads", "Cuff", "Chain", "Pendant", "Pendant Necklace", "Short Necklace", "Mala",
               "Polki Ring", "Pair Bangles"]),
    ("Colour", ["Gold", "Red", "Green", "Blue", "White", "Purple", "Black", "Red Green", "Multi"]),
    ("Polish", ["Gold", "Rose Gold", "Matte Gold", "Rhodium", "Silver Oxidised"]),
    ("Size", ["Free Size", "Adjustable", "2x4", "2x6", "2x8", "2x10"]),
    ("Stone Type", ["CZ", "AD", "Kundan", "Polki", "Pearl", "Glass", "Ruby", "Emerald",
                    "Synthetic", "Amethyst", "Amber"]),
    ("Stone Colour", ["White", "Red", "Green", "Blue", "Pink", "Purple", "Black", "Yellow", "Multi"]),
]
LIST_COL = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(LISTS)}
LIST_MAX = 300
VENDOR_NAMES = f"Vendors!$B${FIRST}:$B${FIRST + N_VENDORS - 1}"

# A column's drop-down: (source range, strict?, message)
DROPDOWNS = {
    "Vendor": (f"={VENDOR_NAMES}", True,
               "Pick a vendor from the list. A new vendor: add a row on the Vendors sheet first (name and state)."),
    "Store": (f"=Lists!${LIST_COL['Store']}$2:${LIST_COL['Store']}${LIST_MAX}", True,
              "Pick a store from the list. A new store must be set up in Wholeops first: tell us."),
    "HSN": (f"={HSN_CODES}", True, "Pick the HSN Code from the list. The Codes sheet says which one fits."),
    "State": (f"={STATE_NAMES}", True, "Pick the state from the list."),
}
for name, _ in LISTS:
    if name != "Store":
        lc = LIST_COL[name]
        DROPDOWNS[name] = (f"=Lists!${lc}$2:${lc}${LIST_MAX}", False,
                           "This value is not on the Lists sheet. Click Yes to keep it: Wholeops creates it on upload. "
                           "Add it to the Lists sheet so the drop-down offers it next time.")

# ---- Columns -------------------------------------------------------------------
# (header, kind, dropdown, width, hint, header colour); kind: req / opt / grn / auto
COLS = [
    ("Check", "auto", None, 30, "# Automatic: says OK, or what to fix", H_AUTO),
    ("PO Ref *", "req", None, 14, "Required. Your name for this PO, e.g. SIY-OCT-01. Same PO Ref = one PO", H_PO),
    ("Vendor *", "req", "Vendor", 14, "Required. Pick from the list (Vendors sheet)", H_PO),
    ("Store *", "req", "Store", 13, "Required. Where the goods arrive", H_PO),
    ("Design No", "opt", None, 11, "Optional. Your label: same label = one design (1, 1, 1). Blank = own design", H_PROD),
    ("Product Type *", "req", "Product Type", 14, "Required. What it is: Earring, Necklace...", H_PROD),
    ("Category *", "req", "Category", 12, "Required. Collection: Antique, Kundan...", H_PROD),
    ("HSN Code *", "req", "HSN", 30, "Required. Pick from the list (see the Codes sheet)", H_TAX),
    ("GST %", "auto", None, 8, "Auto, from the HSN", H_AUTO),
    ("Style", "opt", "Style", 14, "Optional. Ball Drop, Cuff...", H_PROD),
    ("Colour", "opt", "Colour", 11, "Optional. Blank if it does not apply", H_PROD),
    ("Polish", "opt", "Polish", 12, "Optional", H_PROD),
    ("Size", "opt", "Size", 10, "Optional", H_PROD),
    ("Stone Type", "opt", "Stone Type", 11, "Optional", H_PROD),
    ("Stone Colour", "opt", "Stone Colour", 12, "Optional", H_PROD),
    ("Photo", "opt", None, 22, "Optional. Photo file name(s), e.g. EI001-red.jpg, or a.jpg, b.jpg", H_PROD),
    ("Qty *", "req", None, 7, "Required. Pieces (a pair = 1)", H_QTY),
    ("Purchase Rate (₹) *", "req", None, 12, "Required. Per piece, including GST, as on the bill", H_QTY),
    ("MRP (₹) *", "req", None, 10, "Required. The price printed on the sticker", H_QTY),
    ("Amount (₹)", "auto", None, 12, "Automatic: Qty × Rate", H_AUTO),
    ("Vendor Invoice No", "grn", None, 17, "When goods arrive: the bill no. Blank = PO only", H_GRN),
    ("Received Qty", "grn", None, 10, "Only if it differs from Qty (short or extra)", H_GRN),
    ("Received Amount (₹)", "auto", None, 13, "Automatic: received qty × rate", H_AUTO),
]
BILL_COLS = [
    ("Check", "auto", None, 40, "# Automatic: says OK, or what to fix", H_AUTO),
    ("Vendor Invoice No *", "req", None, 20, "Required. Exactly as typed on the Purchase sheet", H_GRN),
    ("Vendor *", "req", "Vendor", 16, "Required. Pick from the list", H_GRN),
    ("Invoice Date *", "req", None, 13, "Required. Date on the bill", H_GRN),
    ("Invoice Total (₹) *", "req", None, 15, "Required. Grand total on the bill, including GST", H_GRN),
    ("Rows Total (₹)", "auto", None, 14, "Automatic: sum of Received Amount for this bill", H_AUTO),
    ("Difference (₹)", "auto", None, 13, "Automatic: bill minus rows. Must be within ₹1", H_AUTO),
]
VENDOR_COLS = [
    ("Check", "auto", None, 34, "# Automatic: says OK, or what to fix", H_AUTO),
    ("Vendor Name *", "req", None, 24, "Required. As you want it on the PO", H_PO),
    ("State *", "req", "State", 22, "Required. Where the vendor is. Decides CGST + SGST or IGST", H_PO),
    ("City", "opt", None, 14, "Optional", H_PO),
    ("GSTIN", "opt", None, 18, "Optional. Leave blank if you do not know it", H_PO),
    ("Phone", "opt", None, 14, "Optional", H_PO),
    ("Email", "opt", None, 26, "Optional. POs can be emailed here", H_PO),
]


def key(h):
    return h.replace(" *", "").replace(" (₹)", "")


L = {key(h): get_column_letter(i + 1) for i, (h, *_r) in enumerate(COLS)}
REQUIRED = [key(h) for h, kind, *_r in COLS if kind == "req"]
KIND_FILL = {"req": C_REQ, "opt": C_OPT, "grn": C_GRN, "auto": C_AUTO}


def input_runs(cols, row, absolute_col=True):
    """The row's input cells as contiguous ranges (formula cells left out)."""
    runs, start = [], None
    for i, (h, kind, *_r) in enumerate(cols + [("", "auto")], start=1):
        if kind != "auto" and start is None:
            start = i
        if kind == "auto" and start is not None:
            a, b = get_column_letter(start), get_column_letter(i - 1)
            d = "$" if absolute_col else ""
            runs.append(f"{d}{a}{row}:{d}{b}{row}" if a != b else f"{d}{a}{row}")
            start = None
    return ",".join(runs)


def header(ws, cols, row=1):
    for i, (h, kind, _dd, width, hint, hcol) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = fill(hcol)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        hc = ws.cell(row=row + 1, column=i, value=hint)
        hc.font = font(italic=True, size=8, color="595959")
        hc.fill = fill("F2F2F2")
        hc.alignment = Alignment(vertical="top", wrap_text=True)
        hc.border = BORDER
        if row == 1:
            ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30
    ws.row_dimensions[row + 1].height = 60


def style_rows(ws, cols, first, last):
    for i, (h, kind, *_r) in enumerate(cols, start=1):
        f = fill(KIND_FILL[kind])
        for r in range(first, last + 1):
            c = ws.cell(row=r, column=i)
            c.fill = f
            c.border = BORDER
            c.font = font(color="404040" if kind == "auto" else "000000")


def missing_expr(cols, r):
    return "&".join(f'IF({get_column_letter(i)}{r}="","{key(h)}, ","")'
                    for i, (h, kind, *_x) in enumerate(cols, start=1) if kind == "req")


def purchase_formulas(ws, first, last, bills_sheet, b_first, b_last):
    """Check / GST / Amount / Received Amount for each Purchase row."""
    rng = lambda col: f"${col}${first}:${col}${last}"
    bills = lambda col: f"{bills_sheet}${col}${b_first}:${col}${b_last}"
    for r in range(first, last + 1):
        g = {k: f"{v}{r}" for k, v in L.items()}
        missing = missing_expr(COLS, r)
        design_clash = (
            f'COUNTIFS({rng(L["Vendor"])},{g["Vendor"]},{rng(L["Design No"])},{g["Design No"]},'
            f'{rng(L["Product Type"])},"<>"&{g["Product Type"]})'
            f'+COUNTIFS({rng(L["Vendor"])},{g["Vendor"]},{rng(L["Design No"])},{g["Design No"]},'
            f'{rng(L["Category"])},"<>"&{g["Category"]})'
        )
        hsn_clash = (
            f'COUNTIFS({rng(L["Product Type"])},{g["Product Type"]},{rng(L["Category"])},{g["Category"]},'
            f'{rng(L["HSN Code"])},"<>"&{g["HSN Code"]})'
        )
        check = (
            f'=IF(COUNTA({input_runs(COLS, r, False)})=0,"",'
            f'IF(({missing})<>"","Missing: "&LEFT({missing},LEN({missing})-2),'
            f'IF(COUNTIF({VENDOR_NAMES},{g["Vendor"]})=0,"Add vendor "&{g["Vendor"]}&" on the Vendors sheet",'
            f'IF(ISNA(MATCH({g["HSN Code"]},{HSN_CODES},0)),"Pick the HSN Code from the list",'
            f'IF(AND({g["Design No"]}<>"",{design_clash}>0),"Design "&{g["Design No"]}&" has another Product Type or Category on another row",'
            f'IF({hsn_clash}>0,{g["Product Type"]}&" + "&{g["Category"]}&" has another HSN Code on another row",'
            f'IF(AND({g["Received Qty"]}<>"",{g["Vendor Invoice No"]}=""),"Received Qty needs the Vendor Invoice No",'
            f'IF(AND({g["Vendor Invoice No"]}<>"",COUNTIFS({bills("B")},{g["Vendor Invoice No"]},{bills("C")},{g["Vendor"]})=0),'
            f'"Add bill "&{g["Vendor Invoice No"]}&" of "&{g["Vendor"]}&" on the Bills sheet",'
            f'IF({g["MRP"]}<{g["Purchase Rate"]},"MRP is below the Purchase Rate","OK")))))))))'
        )
        ws[g["Check"]] = check
        ws[g["GST %"]] = f'=IF({g["HSN Code"]}="","",IFERROR(INDEX({HSN_GST},MATCH({g["HSN Code"]},{HSN_CODES},0)),""))'
        ws[g["Amount"]] = f'=IF(OR({g["Qty"]}="",{g["Purchase Rate"]}=""),"",{g["Qty"]}*{g["Purchase Rate"]})'
        ws[g["Received Amount"]] = (
            f'=IF(OR({g["Vendor Invoice No"]}="",{g["Purchase Rate"]}=""),"",'
            f'IF({g["Received Qty"]}="",{g["Qty"]},{g["Received Qty"]})*{g["Purchase Rate"]})'
        )


def bill_formulas(ws, first, last, p_sheet, p_first, p_last):
    p = lambda col: f"{p_sheet}${col}${p_first}:${col}${p_last}"
    for r in range(first, last + 1):
        missing = missing_expr(BILL_COLS, r)
        ws[f"F{r}"] = (f'=IF(B{r}="","",SUMIFS({p(L["Received Amount"])},{p(L["Vendor Invoice No"])},B{r},'
                       f'{p(L["Vendor"])},C{r}))')
        ws[f"G{r}"] = f'=IF(OR(B{r}="",E{r}=""),"",E{r}-F{r})'
        ws[f"A{r}"] = (
            f'=IF(COUNTA(B{r}:E{r})=0,"",IF(({missing})<>"","Missing: "&LEFT({missing},LEN({missing})-2),'
            f'IF(COUNTIFS($B${first}:$B${last},B{r},$C${first}:$C${last},C{r})>1,"Bill "&B{r}&" is listed twice",'
            f'IF(F{r}=0,"No Purchase row has this bill no for this vendor",'
            f'IF(ABS(G{r})>1,"Bill and rows differ by ₹"&TEXT(G{r},"#,##0.00"),"OK")))))'
        )


def vendor_formulas(ws, first, last):
    for r in range(first, last + 1):
        missing = missing_expr(VENDOR_COLS, r)
        code = f'INDEX({STATE_CODES},MATCH(C{r},{STATE_NAMES},0))'
        ws[f"A{r}"] = (
            f'=IF(COUNTA(B{r}:G{r})=0,"",IF(({missing})<>"","Missing: "&LEFT({missing},LEN({missing})-2),'
            f'IF(COUNTIF($B${first}:$B${last},B{r})>1,"Vendor "&B{r}&" is listed twice",'
            f'IF(ISNA(MATCH(C{r},{STATE_NAMES},0)),"Pick the State from the list",'
            f'IF(E{r}="","OK",'
            f'IF(LEN(TRIM(E{r}))<>15,"GSTIN must be 15 characters - or leave it blank",'
            f'IF(LEFT(TRIM(E{r}),2)<>{code},"GSTIN starts with "&LEFT(TRIM(E{r}),2)&" but "&C{r}&" is "&{code},"OK")))))))'
        )


def formats(ws, first, last, colmap):
    for col, fmt in colmap.items():
        for r in range(first, last + 1):
            ws[f"{col}{r}"].number_format = fmt


def check_cf(ws, first, last):
    a = f"A{first}:A{last}"
    ws.conditional_formatting.add(a, FormulaRule(formula=[f'A{first}="OK"'],
                                                 fill=fill("C6EFCE"), font=Font(name=FONT, color="006100", bold=True)))
    ws.conditional_formatting.add(a, FormulaRule(formula=[f'AND(LEN(A{first})>0,A{first}<>"OK")'],
                                                 fill=fill(C_RED), font=Font(name=FONT, color=C_RED_TXT, bold=True)))


def required_cf(ws, cols, first, last):
    """A required box turns red when its row is started and it is empty."""
    started = f"COUNTA({input_runs(cols, first)})>0"
    for i, (h, kind, *_r) in enumerate(cols, start=1):
        if kind != "req":
            continue
        col = get_column_letter(i)
        ws.conditional_formatting.add(
            f"{col}{first}:{col}{last}",
            FormulaRule(formula=[f'AND({started},LEN(TRIM({col}{first}))=0)'], fill=fill(C_RED), stopIfTrue=True))


def grn_cf(ws, first, last):
    """Vendor Invoice No turns red when a Received Qty is typed without it."""
    inv, rq = L["Vendor Invoice No"], L["Received Qty"]
    ws.conditional_formatting.add(
        f"{inv}{first}:{inv}{last}",
        FormulaRule(formula=[f'AND(LEN({rq}{first})>0,LEN(TRIM({inv}{first}))=0)'], fill=fill(C_RED), stopIfTrue=True))


def list_validations(ws, first, last, cols):
    for i, (h, kind, dd, *_r) in enumerate(cols, start=1):
        if not dd:
            continue
        src, strict, msg = DROPDOWNS[dd]
        dv = DataValidation(type="list", formula1=src, allow_blank=True, errorStyle="stop" if strict else "warning",
                            showErrorMessage=True, errorTitle="Not in the list", error=msg)
        col = get_column_letter(i)
        dv.add(f"{col}{first}:{col}{last}")
        ws.add_data_validation(dv)


def number_validations(ws, first, last, whole_cols, money_cols, date_cols=()):
    for col in whole_cols:
        dv = DataValidation(type="whole", operator="between", formula1="0", formula2="1000000",
                            allow_blank=True, showErrorMessage=True, errorTitle="Whole number",
                            error="Type a whole number of pieces, e.g. 3")
        dv.add(f"{col}{first}:{col}{last}")
        ws.add_data_validation(dv)
    for col in money_cols:
        dv = DataValidation(type="decimal", operator="between", formula1="0.01", formula2="100000000",
                            allow_blank=True, showErrorMessage=True, errorTitle="Amount",
                            error="Type an amount in rupees above 0, e.g. 90 or 349.50")
        dv.add(f"{col}{first}:{col}{last}")
        ws.add_data_validation(dv)
    for col in date_cols:
        dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2035,12,31)",
                            allow_blank=True, showErrorMessage=True, errorTitle="Date",
                            error="Type the bill date, e.g. 20-10-2026")
        dv.add(f"{col}{first}:{col}{last}")
        ws.add_data_validation(dv)


def build_purchase_sheet(ws, first, last, bills_sheet, b_first, b_last):
    header(ws, COLS)
    style_rows(ws, COLS, first, last)
    purchase_formulas(ws, first, last, bills_sheet, b_first, b_last)
    text_cols = [L["PO Ref"], L["Design No"], L["Vendor Invoice No"]]
    formats(ws, first, last, {**{c: "@" for c in text_cols},
                              L["GST %"]: '0"%"',
                              L["Qty"]: "0", L["Received Qty"]: "0",
                              L["Purchase Rate"]: "#,##0.00", L["MRP"]: "#,##0.00",
                              L["Amount"]: "#,##0.00", L["Received Amount"]: "#,##0.00"})
    list_validations(ws, first, last, COLS)
    number_validations(ws, first, last, [L["Qty"], L["Received Qty"]], [L["Purchase Rate"], L["MRP"]])
    check_cf(ws, first, last)
    required_cf(ws, COLS, first, last)
    grn_cf(ws, first, last)


def build_bills_sheet(ws, first, last, p_sheet, p_first, p_last, with_header=True):
    if with_header:
        header(ws, BILL_COLS)
    style_rows(ws, BILL_COLS, first, last)
    bill_formulas(ws, first, last, p_sheet, p_first, p_last)
    formats(ws, first, last, {"B": "@", "D": "dd-mmm-yyyy", "E": "#,##0.00", "F": "#,##0.00", "G": "#,##0.00"})
    list_validations(ws, first, last, BILL_COLS)
    number_validations(ws, first, last, [], ["E"], ["D"])
    check_cf(ws, first, last)
    required_cf(ws, BILL_COLS, first, last)


wb = Workbook()

# ---- Start Here --------------------------------------------------------------
st = wb.active
st.title = "Start Here"
st.sheet_view.showGridLines = False
st.column_dimensions["A"].width = 3
st.column_dimensions["B"].width = 24
st.column_dimensions["C"].width = 14
st.column_dimensions["D"].width = 62
st.column_dimensions["E"].width = 22


def put(r, c, v, **kw):
    cell = st.cell(row=r, column=c, value=v)
    cell.font = font(**{k: kw.pop(k) for k in list(kw) if k in ("bold", "size", "color", "italic")})
    cell.alignment = Alignment(vertical="top", wrap_text=kw.pop("wrap", True))
    if "bg" in kw:
        cell.fill = fill(kw.pop("bg"))
    return cell


def title(r, text):
    put(r, 2, text, bold=True, size=12, color="1F3864", wrap=False)


put(1, 2, f"{CLIENT} - Bulk Purchase Upload" if CLIENT else "Bulk Purchase Upload", bold=True, size=16, color="1F3864", wrap=False)
put(2, 2, "Purchase Orders and Goods Receipt from one Excel file", italic=True, color="595959", wrap=False)

r = 4
title(r, "How it works"); r += 1
steps = [
    ("1. List your vendors",
     "On the Vendors sheet, once: the vendor's name and State. City, GSTIN, phone and email only if you have them."),
    ("2. Fill the Purchase sheet",
     "One row for each design in each colour / size you are ordering. Rows with the same PO Ref become one "
     "Purchase Order. New designs are created automatically; you do not need SKU codes or barcodes."),
    ("3. Send the file to us",
     "With the photos, if any. We upload it in Wholeops. Every PO is created in one click, and you can see and print them in the app."),
    ("4. When the parcel and bill arrive",
     "Fill the green columns on the Purchase sheet (Vendor Invoice No, and Received Qty only if it is short or extra) "
     "and add the bill on the Bills sheet. Send the same file again. The goods are received into stock, barcodes "
     "are issued and stickers are ready to print."),
    ("Goods already here?", "Fill everything, green columns and Bills sheet included, in one go."),
]
for a, b in steps:
    put(r, 2, a, bold=True)
    st.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    put(r, 3, b)
    st.row_dimensions[r].height = 15 * max(1, -(-len(b) // 130)) + 3
    r += 1

r += 1
title(r, "Rules"); r += 1
rules = [
    "One row = one design in one colour / size. Rows with the same Design No are one design in different colours or sizes: use any label you like (1, 2, 3 or EI001) - it only groups the rows of this file, and Wholeops gives every design its own number. Leave it blank and the row is a design of its own. To reorder a design already in Wholeops, write its Wholeops design number (e.g. D00010).",
    "Qty is in pieces you sell. A pair of earrings = 1. A set = 1.",
    "Purchase Rate is per piece, including GST, as on the vendor's bill. MRP is the price printed on the sticker.",
    "Photos: write each photo's file name in the Photo column - any name you like (several: a.jpg, b.jpg) - and send the photos with this file. Send JPG photos (PNG, WebP or GIF also work), up to 10 MB each. Each photo goes to its row's item and is renamed after the item's Wholeops number; the first photo of a design also goes to the design.",
    "HSN Code: pick it from the list (the Codes sheet says which one fits). Artificial / fashion jewellery is 7117; real gold or silver is 7113. GST fills itself: all jewellery is 3%.",
    "Vendor GSTIN is optional. The vendor's State is enough to charge the right GST (same state as your store: CGST + SGST; another state: IGST). Only vendors used on the Purchase sheet are created; sample rows left on the Vendors sheet are ignored.",
    "Leave an optional box blank when it does not apply. Do not write NA or -.",
    "A new product type, category, style, colour, polish, size or stone can simply be typed: Wholeops creates it on upload (add it to the Lists sheet so the drop-down offers it next time). A new store must be set up in Wholeops first: tell us.",
    "A required box turns red while it is empty. Column A (Check) must say OK on every row of the Vendors, Purchase and Bills sheets before you send the file.",
    "Do not rename, move or delete the column headers, row 2 or the sheets. Grey columns fill themselves.",
    "Sending the file again is safe: a PO already created (same PO Ref and vendor) is not created twice.",
]
for t in rules:
    st.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    put(r, 2, "•  " + t)
    st.row_dimensions[r].height = 15 * max(1, -(-len(t) // 145)) + 2
    r += 1

r += 1
title(r, "Colours"); r += 1
for bg, label in [(C_REQ, "You fill it: required"), (C_OPT, "You fill it: optional"),
                  (C_GRN, "You fill it when the goods arrive"), (C_AUTO, "Automatic: do not type here"),
                  (C_RED, "Required but still empty: fill it")]:
    c = put(r, 2, "", bg=bg)
    c.border = BORDER
    put(r, 3, label)
    st.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    r += 1

r += 1
title(r, "Your file right now"); r += 1
P, B, V = "Purchase", "Bills", "Vendors"
pa = f"{P}!$A${FIRST}:$A${FIRST + N_ROWS - 1}"
ba = f"{B}!$A${FIRST}:$A${FIRST + N_BILLS - 1}"
va = f"{V}!$A${FIRST}:$A${FIRST + N_VENDORS - 1}"
summary = [
    ("Vendors", f"=SUMPRODUCT(--(LEN({va})>0))", "0"),
    ("Vendors to fix", f'=SUMPRODUCT((LEN({va})>0)*({va}<>"OK"))', "0"),
    ("Purchase rows", f"=SUMPRODUCT(--(LEN({pa})>0))", "0"),
    ("Rows to fix", f'=SUMPRODUCT((LEN({pa})>0)*({pa}<>"OK"))', "0"),
    ("Total Qty", f"=SUM({P}!${L['Qty']}${FIRST}:${L['Qty']}${FIRST + N_ROWS - 1})", "#,##0"),
    ("Total Amount ₹", f"=SUM({P}!${L['Amount']}${FIRST}:${L['Amount']}${FIRST + N_ROWS - 1})", "#,##0.00"),
    ("Bills", f"=SUMPRODUCT(--(LEN({ba})>0))", "0"),
    ("Bills to fix", f'=SUMPRODUCT((LEN({ba})>0)*({ba}<>"OK"))', "0"),
]
srow = r
for label, f, fmt in summary:
    put(r, 2, label)
    c = st.cell(row=r, column=3, value=f)
    c.font = font(bold=True)
    c.number_format = fmt
    c.alignment = Alignment(horizontal="right")
    r += 1
put(r, 2, "Status", bold=True)
c = st.cell(row=r, column=3,
            value=f'=IF(C{srow + 2}=0,"Not started",IF(C{srow + 1}+C{srow + 3}+C{srow + 7}=0,"Ready to send",'
                  f'"Fix the rows marked red in column A"))')
c.font = font(bold=True)
st.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
r += 2

title(r, "Column guide (Purchase sheet)"); r += 1
for i, h in enumerate(["Column", "Required?", "What to write", "Example"]):
    put(r, 2 + i, h, bold=True, color="FFFFFF", bg="1F3864")
r += 1
guide = [
    ("PO Ref", "Yes", "Any name for the order. All rows with the same PO Ref and vendor become one PO.", "SIY-OCT-01"),
    ("Vendor", "Yes", "Who you buy from. Pick from the list; the list is your Vendors sheet.", "Siyyona"),
    ("Store", "Yes", "Where the goods arrive. Pick from the list.", "Main Store"),
    ("Design No", "No", "Your label to group rows: the same label = one design in several colours or sizes. Blank = the row is its own design. Wholeops numbers every design; to reorder one, write its Wholeops number (e.g. D00010).", "1"),
    ("Product Type", "Yes", "What the piece is.", "Earring"),
    ("Category", "Yes", "The collection or make.", "Imported"),
    ("HSN Code", "Yes", "Pick from the list. The same Product Type + Category always takes the same HSN.", "7117 - Imitation / fashion jewellery"),
    ("GST %", "Automatic", "Filled from the HSN Code.", "3%"),
    ("Style", "No", "The shape or style.", "Huggie Hoop"),
    ("Colour", "No", "Main colour of the piece.", "Purple"),
    ("Polish", "No", "Plating or finish.", "Rose Gold"),
    ("Size", "No", "Size, if it has one.", "2x6"),
    ("Stone Type", "No", "The stone used.", "Glass"),
    ("Stone Colour", "No", "Colour of the stone.", "Purple"),
    ("Photo", "No", "The photo's file name, as you saved it (several: a.jpg, b.jpg). Send the photos with the file.", "EI001-red.jpg"),
    ("Qty", "Yes", "Pieces ordered (a pair = 1).", "3"),
    ("Purchase Rate (₹)", "Yes", "Price per piece including GST.", "90"),
    ("MRP (₹)", "Yes", "The price printed on the sticker.", "349"),
    ("Vendor Invoice No", "When goods arrive", "The bill number. Leave blank until the goods arrive.", "SE/2526/PL/3568"),
    ("Received Qty", "Only if different", "Pieces that actually came, when not the same as Qty.", "2"),
    ("Bills sheet", "When goods arrive", "One row per vendor bill: invoice no, vendor, date and the bill's grand total.", "SE/2526/PL/3568, Siyyona, 20-10-2026, 1,000"),
    ("Vendors sheet", "Yes", "Every vendor once: name and State. GSTIN only if you know it.", "Siyyona, Maharashtra"),
]
for row in guide:
    for i, v in enumerate(row):
        c = put(r, 2 + i, v, bold=(i == 0))
        c.border = BORDER
    lines = max(-(-len(row[1]) // 13), -(-len(row[2]) // 66), -(-len(row[3]) // 21), 1)
    st.row_dimensions[r].height = 15 * lines + 1
    r += 1
r += 1
put(r, 2, f"Template v1.7, {date.today().strftime('%d %b %Y')}. The values on the Vendors, Lists and Example sheets are samples: "
          "replace them with yours. The Codes sheet (HSN, states) is fixed.", italic=True, color="7F7F7F")
st.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)

# ---- Vendors / Purchase / Bills -------------------------------------------------
vs = wb.create_sheet("Vendors")
ps = wb.create_sheet("Purchase")
bs = wb.create_sheet("Bills")
header(vs, VENDOR_COLS)
VLAST = FIRST + N_VENDORS - 1
style_rows(vs, VENDOR_COLS, FIRST, VLAST)
vendor_formulas(vs, FIRST, VLAST)
formats(vs, FIRST, VLAST, {"E": "@", "F": "@"})
list_validations(vs, FIRST, VLAST, VENDOR_COLS)
check_cf(vs, FIRST, VLAST)
required_cf(vs, VENDOR_COLS, FIRST, VLAST)
for i, (name, city) in enumerate([("Siyyona", "Mumbai"), ("Thofa", "Mumbai"), ("YJ NX", "Mumbai"), ("Shwetha", "Mumbai"),
                                  ("Piccasso", "Mumbai"), ("Pashwa", "Mumbai"), ("Bittu", "Mumbai")]):
    vs[f"B{FIRST + i}"] = name
    vs[f"C{FIRST + i}"] = "Maharashtra"
    vs[f"D{FIRST + i}"] = city
vs.freeze_panes = "C3"
vs.sheet_properties.tabColor = "1F3864"

build_purchase_sheet(ps, FIRST, FIRST + N_ROWS - 1, "Bills!", FIRST, FIRST + N_BILLS - 1)
build_bills_sheet(bs, FIRST, FIRST + N_BILLS - 1, "Purchase!", FIRST, FIRST + N_ROWS - 1)
ps.freeze_panes = f"{L['Product Type']}3"
bs.freeze_panes = "C3"
ps.sheet_properties.tabColor = "1F3864"
bs.sheet_properties.tabColor = "375623"

# ---- Example (lines on top, a bills block below) ------------------------------
ex = wb.create_sheet("Example")
ex.sheet_properties.tabColor = "BF8F00"
IMIT, BASE = HSN[0][0], HSN[2][0]
EX_ROWS = [
    dict(r="SIY-OCT-01", v="Siyyona", s="Main Store", d="EI001", pt="Earring", c="Imported", h=IMIT, st="Huggie Hoop",
         col="Purple", pol="Gold", sz="Free Size", stt="Glass", stc="Purple", ph="EI001-purple.jpg", q=3, rate=90, mrp=349, inv="SE/2526/PL/3568",
         why="Goods arrived: bill no filled, all 3 came. Its photo is EI001-purple.jpg"),
    dict(r="SIY-OCT-01", v="Siyyona", s="Main Store", d="EI001", pt="Earring", c="Imported", h=IMIT, st="Huggie Hoop",
         col="White", pol="Rose Gold", sz="Free Size", stt="Glass", stc="White", ph="EI001-white.jpg, EI001-white-back.jpg", q=3, rate=90, mrp=349, inv="SE/2526/PL/3568",
         why="Same design EI001 in another colour = a new row. Two photos"),
    dict(r="SIY-OCT-01", v="Siyyona", s="Main Store", d="EI002", pt="Earring", c="Imported", h=IMIT, st="Ball Drop",
         col="White", pol="Gold", sz="Free Size", q=3, rate=90, mrp=325, inv="SE/2526/PL/3568", rq=2,
         why="Only 2 of 3 came: Received Qty = 2"),
    dict(r="SIY-OCT-01", v="Siyyona", s="Main Store", d="ET001", pt="Earring", c="Temple", h=IMIT, st="Star Stud",
         col="Red", pol="Gold", sz="Free Size", stt="Kundan", stc="Red", q=2, rate=140, mrp=449, inv="SE/2526/PL/3568",
         why="Stone boxes filled because it has stones"),
    dict(r="THO-OCT-01", v="Thofa", s="Main Store", d="EA001", pt="Earring", c="Antique", h=IMIT, st="Ball Drop",
         col="Gold", pol="Matte Gold", sz="Free Size", stt="Ruby", stc="Red", q=6, rate=110, mrp=399,
         why="Goods not here yet: green columns blank, so only the PO is made"),
    dict(r="THO-OCT-01", v="Thofa", s="Main Store", pt="Necklace", c="Antique", h=IMIT, st="Short Necklace",
         col="Red Green", pol="Gold", sz="Adjustable", stt="Ruby", stc="Multi", q=4, rate=250, mrp=899,
         why="No Design No: a design of its own, numbered by Wholeops (same PO Ref = same PO)"),
    dict(r="BIT-OCT-01", v="Bittu", s="Main Store", d="BA001", pt="Bangle", c="Antique", h=BASE, st="Pair Bangles",
         col="Gold", pol="Gold", sz="2x6", q=4, rate=180, mrp=599,
         why="Each size is its own row. Brass bangles: the 6-digit 711719 is also fine"),
    dict(r="BIT-OCT-01", v="Bittu", s="Main Store", d="BA001", pt="Bangle", c="Antique", h=BASE, st="Pair Bangles",
         col="Gold", pol="Gold", sz="2x8", q=4, rate=180, mrp=599, why=""),
]
FIELD = {"r": "PO Ref", "v": "Vendor", "s": "Store", "d": "Design No", "pt": "Product Type", "c": "Category",
         "h": "HSN Code", "st": "Style", "col": "Colour", "pol": "Polish", "sz": "Size", "stt": "Stone Type",
         "stc": "Stone Colour", "ph": "Photo", "q": "Qty", "rate": "Purchase Rate", "mrp": "MRP", "inv": "Vendor Invoice No",
         "rq": "Received Qty"}
EX_FIRST, EX_LAST = FIRST, FIRST + len(EX_ROWS) - 1
EXB_HDR = EX_LAST + 3          # bills header row inside the Example sheet
EXB_FIRST = EXB_HDR + 2
EXB_LAST = EXB_FIRST           # one example bill
build_purchase_sheet(ex, EX_FIRST, EX_LAST, "", EXB_FIRST, EXB_LAST)
WHY_COL = len(COLS) + 2
for i, row in enumerate(EX_ROWS):
    rr = EX_FIRST + i
    for k, v in row.items():
        if k in FIELD:
            ex[f"{L[FIELD[k]]}{rr}"] = v
    nc = ex.cell(row=rr, column=WHY_COL, value=row["why"])
    nc.font = font(italic=True, color="7F4F00")
why = ex.cell(row=1, column=WHY_COL, value="Why this row looks like this")
why.font = font(bold=True, color="7F4F00")
ex.column_dimensions[get_column_letter(WHY_COL)].width = 64

t = ex.cell(row=EXB_HDR - 1, column=1, value="Bills (in your file this is the Bills sheet)")
t.font = font(bold=True, size=11, color="375623")
header(ex, BILL_COLS, row=EXB_HDR)
build_bills_sheet(ex, EXB_FIRST, EXB_LAST, "", EX_FIRST, EX_LAST, with_header=False)
ex[f"B{EXB_FIRST}"] = "SE/2526/PL/3568"
ex[f"C{EXB_FIRST}"] = "Siyyona"
ex[f"D{EXB_FIRST}"] = date(2026, 10, 20)
ex[f"E{EXB_FIRST}"] = 1000
nc = ex.cell(row=EXB_FIRST, column=WHY_COL,
             value="Bill total 1,000 = 3×90 + 3×90 + 2×90 + 2×140 (only what came)")
nc.font = font(italic=True, color="7F4F00")
ex.freeze_panes = ps.freeze_panes

# ---- Lists ----------------------------------------------------------------------
ls = wb.create_sheet("Lists")
ls.sheet_properties.tabColor = "7F7F7F"
for i, (name, vals) in enumerate(LISTS, start=1):
    c = ls.cell(row=1, column=i, value=name)
    c.font = font(bold=True, color="FFFFFF"); c.fill = fill("595959"); c.border = BORDER
    c.alignment = Alignment(horizontal="center")
    ls.column_dimensions[get_column_letter(i)].width = 16
    for j, v in enumerate(vals, start=2):
        ls.cell(row=j, column=i, value=v).font = font()
note = ls.cell(row=1, column=len(LISTS) + 2,
               value="Sample values: replace them with yours and add new ones at the bottom of a column. "
                     "Stores must already exist in Wholeops; every other new value is created on upload. "
                     "Vendors are on the Vendors sheet.")
note.font = font(italic=True, color="7F7F7F")
note.alignment = Alignment(wrap_text=True, vertical="top")
ls.column_dimensions[get_column_letter(len(LISTS) + 2)].width = 40
ls.freeze_panes = "A2"

# ---- Codes: HSN and states (fixed; protected so nothing changes by accident) ----
cs = wb.create_sheet("Codes")
cs.sheet_properties.tabColor = "5B3A7A"
for col, (h, w) in zip("ABCDEF", [("HSN Code", 52), ("GST %", 8), ("Use it for", 100), ("", 3),
                                   ("State", 36), ("GST State Code", 10)]):
    cs.column_dimensions[col].width = w
    if h:
        c = cs[f"{col}1"]
        c.value = h
        c.font = font(bold=True, color="FFFFFF"); c.fill = fill(H_TAX); c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
cs.row_dimensions[1].height = 28
for j, (code, gst, use) in enumerate(HSN, start=2):
    cs[f"A{j}"] = code
    cs[f"A{j}"].font = font(bold=len(code.split(" ")[0]) == 4)
    cs[f"B{j}"] = gst
    cs[f"B{j}"].number_format = '0"%"'
    cs[f"C{j}"] = use
    for col in "ABC":
        cs[f"{col}{j}"].border = BORDER
        if col != "A":
            cs[f"{col}{j}"].font = font()
for j, (name, code) in enumerate(STATES, start=2):
    cs[f"E{j}"] = name
    cs[f"F{j}"] = code
    cs[f"F{j}"].number_format = "@"
    for col in "EF":
        cs[f"{col}{j}"].border = BORDER
        cs[f"{col}{j}"].font = font()
nr = len(HSN) + 3
for text in [
    "All chapter 71 jewellery is GST 3% (1.5% CGST + 1.5% SGST, or 3% IGST).",
    "On invoices, 4 HSN digits are enough up to ₹5 crore yearly turnover; above that use 6 digits.",
    "Artificial / fashion jewellery is 7117 (711719 for base metal, 711790 for beads, glass, plastic).",
    "Use 7113 only for real gold, silver or platinum jewellery.",
    "This sheet is fixed and protected. Need another HSN? Ask us.",
]:
    cs[f"A{nr}"] = text
    cs[f"A{nr}"].font = font(italic=True, color="595959")
    nr += 1
cs.freeze_panes = "A2"
cs.protection.sheet = True

wb.active = 0
wb.save(OUT)
print("wrote", OUT, "sheets:", wb.sheetnames)
